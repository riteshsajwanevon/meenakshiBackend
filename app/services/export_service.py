"""Exports a report's fields as XLSX, CSV or PDF.

Columns: Row, Field, Original OCR, Corrected, Value, Confidence, Status.
"""

import csv
import io
from dataclasses import dataclass

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Font

from app.core.errors import BadRequestError
from app.schemas.report import ReportDetail

COLUMNS = ["Row", "Field", "Original OCR", "Corrected", "Value", "Confidence", "Status"]
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# requested format -> (file extension, media type)
FORMATS = {
    "xlsx": ("xlsx", XLSX_MEDIA_TYPE),
    "excel": ("xlsx", XLSX_MEDIA_TYPE),
    "csv": ("csv", "text/csv"),
    "pdf": ("pdf", "application/pdf"),
}


@dataclass(frozen=True)
class ExportedFile:
    data: bytes
    content_type: str
    filename: str


def export_report(report: ReportDetail, requested_format: str) -> ExportedFile:
    key = (requested_format or "xlsx").strip().lower()
    if key not in FORMATS:
        raise BadRequestError("Unsupported export format. Use xlsx, csv or pdf.")
    extension, media_type = FORMATS[key]

    rows = _rows(report)
    builders = {"xlsx": _to_xlsx, "csv": _to_csv, "pdf": _to_pdf}
    return ExportedFile(builders[extension](report, rows), media_type, f"report-{report.id}.{extension}")


def _rows(report: ReportDetail) -> list[list]:
    return [
        [
            field.row_index + 1 if field.row_index is not None else "",
            field.label,
            field.original_value or "",
            field.corrected_value or "",
            field.value or "",
            round(field.confidence, 4) if field.confidence is not None else "",
            field.status.value,
        ]
        for field in report.fields
    ]


def _to_xlsx(report: ReportDetail, rows: list[list]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    sheet.append(COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    for column, width in zip("ABCDEFG", (6, 24, 30, 30, 30, 12, 14)):
        sheet.column_dimensions[column].width = width

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _to_csv(report: ReportDetail, rows: list[list]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(COLUMNS)
    writer.writerows(rows)
    return buffer.getvalue().encode("utf-8-sig")  # BOM so Excel detects UTF-8


def _to_pdf(report: ReportDetail, rows: list[list]) -> bytes:
    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=12)
    pdf.add_page()

    pdf.set_font("Helvetica", style="B", size=14)
    pdf.cell(0, 8, _pdf_text(f"{report.report_type} - {report.document_name}"), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", size=9)
    summary = f"Status: {report.status.value}   Month: {report.report_month or '-'}   Uploaded by: {report.uploaded_by}"
    pdf.cell(0, 6, _pdf_text(summary), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)

    pdf.set_font("Helvetica", size=8)
    with pdf.table(col_widths=(10, 40, 50, 50, 50, 20, 25), text_align="LEFT", first_row_as_headings=True) as table:
        for values in [COLUMNS, *rows]:
            table_row = table.row()
            for value in values:
                table_row.cell(_pdf_text(str(value)))
    return bytes(pdf.output())


def _pdf_text(text: str) -> str:
    """The built-in PDF fonts only support Latin-1; unsupported characters become '?'."""
    return text.encode("latin-1", "replace").decode("latin-1")
