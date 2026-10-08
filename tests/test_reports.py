"""Report workflow: upload -> process (fake OCR) -> correct -> validate -> approve, plus files, exports and dashboard."""

import io
import json

from fpdf import FPDF
from openpyxl import load_workbook

from app.services.ocr_client import UNAVAILABLE_USER_MESSAGE, OcrServiceError

from tests.conftest import auth_headers, template_id, upload_png


def _field(detail: dict, key: str, row: int | None = None) -> dict:
    return next(f for f in detail["fields"] if f["fieldKey"] == key and f["rowIndex"] == row)


def _process(client, headers, report_id) -> dict:
    response = client.post(f"/api/v1/reports/{report_id}/process", headers=headers)
    assert response.status_code == 200, response.text
    return response.json()


# ---------------------------------------------------------------- upload

def test_upload_creates_report_and_preview(client, operator_headers):
    summary = upload_png(client, operator_headers, name="June scan.png")

    assert summary["status"] == "UPLOADED"
    assert summary["documentName"] == "June scan.png"
    assert summary["reportType"] == "Supplier Rejection Report"
    assert summary["templateCode"] == "SUPPLIER_REJECTION"
    assert summary["uploadedBy"] == "Line Operator"
    assert summary["uploadedAt"].endswith("Z")

    preview = client.get(f"/api/v1/reports/{summary['id']}/preview", headers=operator_headers)
    assert preview.status_code == 200
    assert preview.headers["content-type"] == "image/png"
    original = client.get(f"/api/v1/reports/{summary['id']}/file", headers=operator_headers)
    assert original.headers["content-disposition"].startswith('inline; filename="June scan.png"')


def test_upload_pdf_renders_preview(client, operator_headers):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(0, 10, "Supplier Rejection Report")
    response = client.post(
        "/api/v1/reports",
        headers=operator_headers,
        data={"templateId": str(template_id(client, operator_headers, "SUPPLIER_REJECTION")), "departmentId": ""},
        files={"file": ("report.pdf", bytes(pdf.output()), "application/pdf")},
    )
    assert response.status_code == 200, response.text
    preview = client.get(f"/api/v1/reports/{response.json()['id']}/preview", headers=operator_headers)
    assert preview.content.startswith(b"\x89PNG")


def test_upload_rejects_bad_files(client, operator_headers):
    template = str(template_id(client, operator_headers, "SUPPLIER_REJECTION"))

    def upload(name, data, content_type):
        return client.post("/api/v1/reports", headers=operator_headers, data={"templateId": template}, files={"file": (name, data, content_type)})

    assert upload("notes.txt", b"hello", "text/plain").status_code == 400
    assert upload("empty.png", b"", "image/png").status_code == 400
    corrupt = upload("broken.png", b"not really a png", "image/png")
    assert corrupt.status_code == 400
    assert corrupt.json()["message"] == "We could not open this file. Please upload a clear PDF or image."


def test_viewer_cannot_upload(client, admin_headers):
    viewer = client.post(
        "/api/v1/users", headers=admin_headers, json={"email": "viewer@example.com", "fullName": "Viewer", "role": "VIEWER"}
    ).json()
    headers = auth_headers(client, ("viewer@example.com", viewer["temporaryPassword"]))
    response = client.post("/api/v1/reports", headers=headers, data={"templateId": "1"}, files={"file": ("a.png", b"x", "image/png")})
    assert response.status_code == 403


# ---------------------------------------------------------------- processing

def test_process_maps_ocr_fields_and_statuses(client, operator_headers, fake_ocr):
    report = upload_png(client, operator_headers)
    detail = _process(client, operator_headers, report["id"])

    assert detail["status"] == "VALIDATION_REQUIRED"
    assert detail["reportMonth"] == "June 2026"
    assert detail["reportDate"] == "2026-06-02"
    assert detail["previewUrl"] == f"/api/v1/reports/{report['id']}/preview"
    assert 0 < detail["averageConfidence"] < 1

    assert _field(detail, "month")["status"] == "VERIFIED"
    assert _field(detail, "partName", 0)["status"] == "VERIFIED"
    assert _field(detail, "rejectionQty", 0)["status"] == "NEEDS_REVIEW"
    assert _field(detail, "problemDescription", 0)["status"] == "ERROR"
    assert _field(detail, "remarks", 0)["status"] == "EMPTY"
    assert _field(detail, "rejectionQty", 1)["status"] == "ERROR"
    assert _field(detail, "lotQty", 0)["boundingBox"] == {"x": 0.1, "y": 0.2, "width": 0.1, "height": 0.03}
    assert not any(f["fieldKey"] == "unknownColumn" for f in detail["fields"])

    # Header first, then rows in order.
    assert detail["fields"][0]["fieldKey"] == "month"

    # The OCR service received the template payload.
    sent = fake_ocr.calls[0]["template"].model_dump(by_alias=True)
    assert sent["code"] == "SUPPLIER_REJECTION"
    assert sent["fields"][0]["fieldKey"] == "month"

    extractions = client.get(f"/api/v1/reports/{report['id']}/extractions", headers=operator_headers).json()
    assert json.loads(extractions[0])["engine"] == "paddleocr"

    notifications = client.get("/api/v1/notifications", headers=operator_headers).json()
    assert notifications[0]["title"] == "Report ready for review"
    assert notifications[0]["read"] is False
    assert client.post(f"/api/v1/notifications/{notifications[0]['id']}/read", headers=operator_headers).status_code == 200
    assert client.get("/api/v1/notifications", headers=operator_headers).json()[0]["read"] is True


def test_ocr_failure_marks_report_failed_and_allows_retry(client, operator_headers, fake_ocr):
    report = upload_png(client, operator_headers)
    fake_ocr.error = OcrServiceError("Connection refused", UNAVAILABLE_USER_MESSAGE)

    response = client.post(f"/api/v1/reports/{report['id']}/process", headers=operator_headers)
    assert response.status_code == 400
    assert response.json()["message"] == UNAVAILABLE_USER_MESSAGE

    failed = client.get(f"/api/v1/reports/{report['id']}", headers=operator_headers).json()
    assert failed["status"] == "FAILED"
    assert failed["failureReason"] == "Connection refused"

    fake_ocr.error = None
    assert _process(client, operator_headers, report["id"])["status"] == "VALIDATION_REQUIRED"


def test_processed_report_cannot_be_processed_again(client, operator_headers, fake_ocr):
    report = upload_png(client, operator_headers)
    _process(client, operator_headers, report["id"])
    assert client.post(f"/api/v1/reports/{report['id']}/process", headers=operator_headers).status_code == 409


# ---------------------------------------------------------------- review workflow

def test_full_review_workflow(client, operator_headers, inspector_headers, fake_ocr):
    report_id = upload_png(client, operator_headers)["id"]
    detail = _process(client, operator_headers, report_id)

    # Operators cannot correct fields (inspector/admin only).
    bad_qty = _field(detail, "rejectionQty", 1)
    url = f"/api/v1/reports/{report_id}/fields/{bad_qty['id']}"
    assert client.patch(url, headers=operator_headers, json={"value": "100"}).status_code == 403

    # Validation fails while a value is invalid, listing the problem.
    blocked = client.post(f"/api/v1/reports/{report_id}/validate", headers=inspector_headers)
    assert blocked.status_code == 400
    assert "Row 2: Rejection Qty must be a number" in blocked.json()["details"]

    # Invalid corrections are rejected; valid ones are saved.
    assert client.patch(url, headers=inspector_headers, json={"value": "-5"}).status_code == 400
    corrected = client.patch(url, headers=inspector_headers, json={"value": "1,00"})
    assert corrected.status_code == 200
    body = corrected.json()
    assert (body["originalValue"], body["correctedValue"], body["value"], body["status"]) == ("abc", "1,00", "1,00", "VERIFIED")
    assert body["correctedBy"] == "Quality Inspector"

    validated = client.post(f"/api/v1/reports/{report_id}/validate", headers=inspector_headers)
    assert validated.status_code == 200, validated.text
    assert validated.json()["status"] == "VALIDATED"

    # Correcting a validated report sends it back for validation.
    part = _field(validated.json(), "partName", 0)
    client.patch(f"/api/v1/reports/{report_id}/fields/{part['id']}", headers=inspector_headers, json={"value": "Pipe 12.70 X 1.5"})
    assert client.get(f"/api/v1/reports/{report_id}", headers=inspector_headers).json()["status"] == "VALIDATION_REQUIRED"

    assert client.post(f"/api/v1/reports/{report_id}/approve", headers=inspector_headers).status_code == 409
    client.post(f"/api/v1/reports/{report_id}/validate", headers=inspector_headers)
    approved = client.post(f"/api/v1/reports/{report_id}/approve", headers=inspector_headers)
    assert approved.status_code == 200
    assert approved.json()["status"] == "APPROVED"
    assert approved.json()["approvedBy"] == "Quality Inspector"

    # Approved reports are locked.
    locked = client.patch(f"/api/v1/reports/{report_id}/fields/{part['id']}", headers=inspector_headers, json={"value": "x"})
    assert locked.status_code == 409


def test_operator_sees_only_own_reports(client, admin_headers, operator_headers, fake_ocr):
    admin_report = upload_png(client, admin_headers, name="admin.png")
    own_report = upload_png(client, operator_headers, name="mine.png")

    listing = client.get("/api/v1/reports", headers=operator_headers).json()
    assert [r["id"] for r in listing["content"]] == [own_report["id"]]

    forbidden = client.get(f"/api/v1/reports/{admin_report['id']}", headers=operator_headers)
    assert forbidden.status_code == 403
    assert forbidden.json()["message"] == "You can only review reports you uploaded"
    assert client.post(f"/api/v1/reports/{admin_report['id']}/process", headers=operator_headers).status_code == 403

    assert client.get("/api/v1/reports", headers=admin_headers).json()["totalElements"] == 2


def test_list_filters_and_pagination(client, operator_headers, inspector_headers, fake_ocr):
    first = upload_png(client, operator_headers, name="first.png")
    upload_png(client, operator_headers, name="second.png")
    _process(client, operator_headers, first["id"])

    def ids(**params):
        return [r["id"] for r in client.get("/api/v1/reports", headers=inspector_headers, params=params).json()["content"]]

    assert ids(status="VALIDATION_REQUIRED") == [first["id"]]
    assert ids(supplier="uttam") == [first["id"]]
    assert ids(part="MIG wire") == [first["id"]]
    assert ids(problem="RUSTY") == [first["id"]]
    assert ids(q="second") != [first["id"]]
    assert ids(q="June 2026") == [first["id"]]
    assert ids(supplier="100%") == []  # wildcards in input are treated literally

    page = client.get("/api/v1/reports", headers=inspector_headers, params={"size": 1, "page": 1}).json()
    assert (page["page"], page["size"], page["totalElements"], page["totalPages"]) == (1, 1, 2, 2)
    assert len(page["content"]) == 1


# ---------------------------------------------------------------- export & dashboard

def test_exports(client, operator_headers, fake_ocr):
    report_id = upload_png(client, operator_headers)["id"]
    _process(client, operator_headers, report_id)

    xlsx = client.get(f"/api/v1/reports/{report_id}/export", headers=operator_headers)
    assert xlsx.status_code == 200
    assert xlsx.headers["content-disposition"].startswith('attachment; filename="report-')
    sheet = load_workbook(io.BytesIO(xlsx.content)).active
    assert [cell.value for cell in sheet[1]] == ["Row", "Field", "Original OCR", "Corrected", "Value", "Confidence", "Status"]

    csv = client.get(f"/api/v1/reports/{report_id}/export", headers=operator_headers, params={"format": "csv"})
    assert csv.headers["content-type"].startswith("text/csv")
    assert "Rejection Qty" in csv.content.decode("utf-8-sig")

    pdf = client.get(f"/api/v1/reports/{report_id}/export", headers=operator_headers, params={"format": "pdf"})
    assert pdf.content.startswith(b"%PDF")

    assert client.get(f"/api/v1/reports/{report_id}/export", headers=operator_headers, params={"format": "doc"}).status_code == 400


def test_dashboard(client, operator_headers, fake_ocr):
    report_id = upload_png(client, operator_headers)["id"]
    upload_png(client, operator_headers)  # stays UPLOADED
    _process(client, operator_headers, report_id)

    summary = client.get("/api/v1/dashboard/summary", headers=operator_headers).json()
    assert summary["totalReports"] == 2
    assert summary["pendingValidation"] == 1
    assert summary["rejectionQuantity"] == 250.0  # row 2 ("abc") is not a number yet
    assert summary["reworkQuantity"] == 0.0

    rejections = client.get("/api/v1/dashboard/rejections", headers=operator_headers).json()
    assert rejections[0]["name"].lower() == "rusty"
    assert rejections[0]["count"] == 2  # "Rusty" and "rusty" grouped together
    assert rejections[0]["quantity"] == 250.0

    suppliers = client.get("/api/v1/dashboard/suppliers", headers=operator_headers).json()
    assert {s["name"] for s in suppliers} == {"UTTAM", "MW Wire Tech"}

    trends = client.get("/api/v1/dashboard/trends", headers=operator_headers, params={"granularity": "monthly"}).json()
    assert trends[0]["reports"] == 2
    assert trends[0]["quantity"] == 250.0
    assert client.get("/api/v1/dashboard/trends", headers=operator_headers, params={"granularity": "hourly"}).status_code == 400
